"""Form setup for calls and tickets: fields, value lists, the scenario tree, and copying between clients."""
import csv
import io
import re
from typing import Any, Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..database import get_db
from ..deps import audit, get_current_user, permissions_of
from ..models import Campaign, Client, FormField, Scenario, ScenarioField, User, ValueItem, ValueList
from ..services import forms as fs

router = APIRouter(tags=["forms"])


# ---------- access ----------

def _staff(u: User) -> bool:
    return u.user_type == "staff"


def scope(db: Session, user: User, client_id: int, context: Optional[str] = None, write: bool = False) -> Client:
    """Checks the person may see (or change) this client's forms. Client users are limited to their own company."""
    perms = permissions_of(user)
    if write and "forms.manage" not in perms:
        raise HTTPException(403, "You don't have permission to change forms")
    if not write and not perms & {"forms.view", "forms.manage"}:
        raise HTTPException(403, "You don't have permission for this")
    client = db.get(Client, client_id)
    if not client or (not _staff(user) and client.id != user.client_id):
        raise HTTPException(404, "Client not found")  # same answer for other companies' ids
    if context and context not in fs.CONTEXTS:
        raise HTTPException(404, "Unknown form")
    return client


def form_scope(db: Session, user: User, client_id: int, key: str, write: bool = False) -> tuple[Client, str, Optional[int]]:
    """Which form is meant. `key` is "in_call" (one per client) or "campaign-<id>" (the out-call form of that campaign).
    Returns (client, context, campaign_id)."""
    client = scope(db, user, client_id, None, write)
    if key == "in_call":
        return client, "in_call", None
    m = re.fullmatch(r"campaign-(\d+)", key)
    campaign = db.get(Campaign, int(m.group(1))) if m else None
    if not campaign or campaign.client_id != client.id:
        raise HTTPException(404, "Form not found")
    return client, "out_call", campaign.id


def fail(fields: dict[str, str], message: str = "Please fix the highlighted fields") -> None:
    raise HTTPException(422, detail={"message": message, "fields": fields})


# ---------- bodies ----------

class FieldIn(BaseModel):
    label: str = Field(default="", max_length=120)
    field_type: str = "text"
    validation: str = "none"
    rules: dict[str, Any] = {}
    required: bool = False
    help_text: Optional[str] = Field(default=None, max_length=250)
    show_always: bool = True
    is_active: bool = True
    value_list_id: Optional[int] = None
    list_level: int = 1
    parent_field_id: Optional[int] = None
    own_values: list[str] = []  # starting values for a dropdown that has its own list


class ReorderIn(BaseModel):
    ids: list[int]


class ScenarioIn(BaseModel):
    name: str = Field(min_length=1, max_length=150)
    parent_id: Optional[int] = None
    outcome: Optional[str] = None
    is_active: bool = True


class LinksIn(BaseModel):
    links: list[dict[str, Any]]  # [{"field_id": 3, "required": true}]


class MoveIn(BaseModel):
    direction: str = Field(pattern="^(up|down)$")


class ListIn(BaseModel):
    client_id: Optional[int] = None
    name: str = Field(min_length=2, max_length=120)
    description: Optional[str] = Field(default=None, max_length=250)
    level_names: list[str] = Field(min_length=1, max_length=fs.MAX_LIST_LEVELS)
    is_active: bool = True


class ItemIn(BaseModel):
    label: str = Field(min_length=1, max_length=200)
    parent_id: Optional[int] = None
    is_active: bool = True


class CheckIn(BaseModel):
    path: list[int] = []
    answers: dict[str, Any] = {}


class CopyIn(BaseModel):
    source_client_id: int
    source_campaign_id: Optional[int] = None  # for a campaign form: which campaign to copy from
    replace: bool = False


# ---------- the whole form of one client ----------

@router.get("/api/forms/clients")
def clients_for_forms(user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    if not permissions_of(user) & {"forms.view", "forms.manage"}:
        raise HTTPException(403, "You don't have permission for this")
    stmt = select(Client).where(Client.status.in_(("active", "suspended"))).order_by(Client.company_name)
    if not _staff(user):
        stmt = stmt.where(Client.id == user.client_id)
    return [{"id": c.id, "name": c.company_name, "code": c.code} for c in db.scalars(stmt)]


@router.get("/api/forms/{client_id}/{context}")
def get_form(client_id: int, context: str, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    client, ctx, cid = form_scope(db, user, client_id, context)
    campaign = db.get(Campaign, cid) if cid else None
    return {**fs.build_config(db, client, ctx, cid, campaign), "can_manage": "forms.manage" in permissions_of(user)}


# ---------- fields ----------

def _list_for(db: Session, user: User, client_id: int, list_id: int) -> ValueList:
    vl = db.get(ValueList, list_id)
    if not vl or (vl.client_id not in (None, client_id)):
        fail({"value_list_id": "Choose one of the available lists"})
    return vl


def _validate_field(db: Session, user: User, client_id: int, context: str, body: FieldIn, existing: Optional[FormField],
                    campaign_id: Optional[int] = None) -> Optional[ValueList]:
    errors = fs.check_field_def(body.model_dump())
    vl = None
    if body.field_type in fs.CHOICE_TYPES:
        if body.value_list_id:
            vl = _list_for(db, user, client_id, body.value_list_id)
            if not 1 <= body.list_level <= len(vl.level_names):
                errors["list_level"] = "This list does not have that level"
            elif body.list_level > 1 and not body.parent_field_id:
                errors["parent_field_id"] = (f"Level {body.list_level} values depend on the level above: "
                                             "choose the dropdown this one depends on")
            elif body.list_level == 1 and body.parent_field_id:
                errors["list_level"] = "A dependent dropdown must use level 2 or deeper of the list"
        if body.parent_field_id:
            parent = db.get(FormField, body.parent_field_id)
            if (not parent or parent.client_id != client_id or parent.context != context or parent.campaign_id != campaign_id
                    or parent.field_type not in ("select", "radio") or (existing and parent.id == existing.id)):
                errors["parent_field_id"] = "Choose a dropdown from this form"
            elif existing and _depends_on(db, parent, existing.id):
                errors["parent_field_id"] = "That field already depends on this one"
            elif not vl or parent.value_list_id != vl.id or body.list_level != parent.list_level + 1:
                errors["parent_field_id"] = ("A dependent dropdown must use the next level of the same list as the field above "
                                             f"(level {parent.list_level + 1} of its list)")
    elif body.parent_field_id or body.value_list_id:
        errors["field_type"] = "Only dropdowns, multi-selects and radio buttons use values"
    if errors:
        fail(errors)
    return vl


def _depends_on(db: Session, field: FormField, target_id: int) -> bool:
    seen = set()
    while field and field.parent_field_id and field.id not in seen:
        if field.parent_field_id == target_id:
            return True
        seen.add(field.id)
        field = db.get(FormField, field.parent_field_id)
    return False


def _free_key(db: Session, client_id: int, context: str, label: str, campaign_id: Optional[int] = None) -> str:
    base = fs.slug(label)
    if base in fs.RESERVED_KEYS:
        base = f"{base}_field"
    key, n = base, 1
    while db.scalar(select(FormField.id).where(FormField.client_id == client_id, FormField.context == context,
                                               fs.in_scope(FormField, campaign_id), FormField.key == key)):
        n += 1
        key = f"{base[:36]}_{n}"
    return key


@router.post("/api/forms/{client_id}/{context}/fields", status_code=201)
def create_field(client_id: int, context: str, body: FieldIn, user: User = Depends(get_current_user),
                 db: Session = Depends(get_db)):
    client, ctx, cid = form_scope(db, user, client_id, context, write=True)
    vl = _validate_field(db, user, client.id, ctx, body, None, campaign_id=cid)
    f = FormField(client_id=client.id, context=ctx, campaign_id=cid,
                  key=_free_key(db, client.id, ctx, body.label.strip(), cid),
                  label=body.label.strip(), field_type=body.field_type, validation=body.validation, rules=body.rules,
                  required=body.required, help_text=body.help_text, show_always=body.show_always, is_active=body.is_active,
                  value_list_id=vl.id if vl else None, list_level=body.list_level if vl else 1,
                  parent_field_id=body.parent_field_id,
                  sort_order=(db.scalar(select(func.max(FormField.sort_order)).where(
                      FormField.client_id == client.id, FormField.context == ctx, fs.in_scope(FormField, cid))) or 0) + 1)
    if body.field_type in fs.CHOICE_TYPES and not vl:
        # a dropdown with its own values: a private one-level list, edited from the field's "Values" button
        own = ValueList(client_id=client.id, name=body.label.strip(), level_names=[body.label.strip()[:120]], is_private=True)
        db.add(own)
        db.flush()
        fs.import_rows(db, own, [[v] for v in body.own_values if v.strip()])
        f.value_list_id, f.list_level = own.id, 1
    db.add(f)
    db.flush()
    audit(db, user, "form_field_created", client_id=client.id, detail={"field": f.id, "label": f.label})
    db.commit()
    return fs.field_out(f)


def _field_for(db: Session, user: User, field_id: int, write: bool) -> FormField:
    f = db.get(FormField, field_id)
    if not f:
        raise HTTPException(404, "Field not found")
    scope(db, user, f.client_id, f.context, write=write)
    return f


@router.put("/api/form-fields/{field_id}")
def update_field(field_id: int, body: FieldIn, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    f = _field_for(db, user, field_id, True)
    if body.field_type != f.field_type:
        fail({"field_type": "The type cannot be changed after the field is created. Add a new field instead."})
    vl = _validate_field(db, user, f.client_id, f.context, body, f, campaign_id=f.campaign_id)
    f.label, f.validation, f.rules = body.label.strip(), body.validation, body.rules
    f.required, f.help_text, f.show_always, f.is_active = body.required, body.help_text, body.show_always, body.is_active
    if f.field_type in fs.CHOICE_TYPES:
        if vl:
            f.value_list_id, f.list_level = vl.id, body.list_level
        f.parent_field_id = body.parent_field_id
    audit(db, user, "form_field_updated", client_id=f.client_id, detail={"field": f.id})
    db.commit()
    return fs.field_out(f)


@router.delete("/api/form-fields/{field_id}")
def delete_field(field_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    f = _field_for(db, user, field_id, True)
    child = db.scalar(select(FormField.label).where(FormField.parent_field_id == f.id))
    if child:
        raise HTTPException(409, f"“{child}” depends on this field. Remove or change that field first.")
    own = db.get(ValueList, f.value_list_id) if f.value_list_id else None
    audit(db, user, "form_field_deleted", client_id=f.client_id, detail={"label": f.label})
    db.delete(f)
    db.flush()
    if own and own.is_private:
        db.delete(own)
    db.commit()
    return {"ok": True}


@router.post("/api/forms/{client_id}/{context}/fields/reorder")
def reorder_fields(client_id: int, context: str, body: ReorderIn, user: User = Depends(get_current_user),
                   db: Session = Depends(get_db)):
    _, ctx, cid = form_scope(db, user, client_id, context, write=True)
    fields = {f.id: f for f in fs.form_fields(db, client_id, ctx, campaign_id=cid)}
    for pos, fid in enumerate(i for i in body.ids if i in fields):
        fields[fid].sort_order = pos + 1
    db.commit()
    return {"ok": True}


# ---------- scenarios ----------

def _scenario_for(db: Session, user: User, sid: int, write: bool) -> Scenario:
    s = db.get(Scenario, sid)
    if not s:
        raise HTTPException(404, "Scenario not found")
    scope(db, user, s.client_id, s.context, write=write)
    return s


def _sibling_clash(db: Session, client_id: int, context: str, parent_id: Optional[int], name: str, skip: Optional[int] = None,
                   campaign_id: Optional[int] = None) -> None:
    stmt = select(Scenario.id).where(Scenario.client_id == client_id, Scenario.context == context,
                                     fs.in_scope(Scenario, campaign_id), func.lower(Scenario.name) == name.lower())
    stmt = stmt.where(Scenario.parent_id.is_(None) if parent_id is None else Scenario.parent_id == parent_id)
    if skip:
        stmt = stmt.where(Scenario.id != skip)
    if db.scalar(stmt):
        fail({"name": "This option already exists at this level"})


@router.post("/api/forms/{client_id}/{context}/scenarios", status_code=201)
def create_scenario(client_id: int, context: str, body: ScenarioIn, user: User = Depends(get_current_user),
                    db: Session = Depends(get_db)):
    client, ctx, cid = form_scope(db, user, client_id, context, write=True)
    level, parent = 1, None
    if body.parent_id:
        parent = db.get(Scenario, body.parent_id)
        if not parent or parent.client_id != client.id or parent.context != ctx or parent.campaign_id != cid:
            raise HTTPException(404, "Parent not found")
        level = parent.level + 1
        if level > fs.MAX_SCENARIO_LEVELS:
            fail({"name": f"The tree can have at most {fs.MAX_SCENARIO_LEVELS} levels"})
    if body.outcome and body.outcome not in fs.OUTCOMES:
        fail({"outcome": "Choose a valid outcome"})
    name = " ".join(body.name.split())
    _sibling_clash(db, client.id, ctx, body.parent_id, name, None, cid)
    last = db.scalar(select(func.max(Scenario.sort_order)).where(
        Scenario.client_id == client.id, Scenario.context == ctx, fs.in_scope(Scenario, cid),
        Scenario.parent_id.is_(None) if not body.parent_id else Scenario.parent_id == body.parent_id)) or 0
    s = Scenario(client_id=client.id, context=ctx, campaign_id=cid, parent_id=body.parent_id, level=level, name=name,
                 outcome=body.outcome or None, sort_order=last + 1, is_active=body.is_active)
    db.add(s)
    db.flush()
    audit(db, user, "scenario_created", client_id=client.id, detail={"scenario": s.id, "name": name})
    db.commit()
    return {"id": s.id, "level": s.level}


@router.put("/api/scenarios/{scenario_id}")
def update_scenario(scenario_id: int, body: ScenarioIn, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    s = _scenario_for(db, user, scenario_id, True)
    if body.outcome and body.outcome not in fs.OUTCOMES:
        fail({"outcome": "Choose a valid outcome"})
    name = " ".join(body.name.split())
    _sibling_clash(db, s.client_id, s.context, s.parent_id, name, s.id, s.campaign_id)
    s.name, s.outcome, s.is_active = name, body.outcome or None, body.is_active
    audit(db, user, "scenario_updated", client_id=s.client_id, detail={"scenario": s.id})
    db.commit()
    return {"ok": True}


@router.delete("/api/scenarios/{scenario_id}")
def delete_scenario(scenario_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    s = _scenario_for(db, user, scenario_id, True)
    audit(db, user, "scenario_deleted", client_id=s.client_id, detail={"name": s.name})
    db.delete(s)  # everything below it goes too
    db.commit()
    return {"ok": True}


@router.post("/api/scenarios/{scenario_id}/move")
def move_scenario(scenario_id: int, body: MoveIn, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    s = _scenario_for(db, user, scenario_id, True)
    stmt = select(Scenario).where(Scenario.client_id == s.client_id, Scenario.context == s.context,
                                  fs.in_scope(Scenario, s.campaign_id))
    stmt = stmt.where(Scenario.parent_id.is_(None) if s.parent_id is None else Scenario.parent_id == s.parent_id)
    sibs = list(db.scalars(stmt.order_by(Scenario.sort_order, Scenario.id)))
    i = sibs.index(s)
    j = i - 1 if body.direction == "up" else i + 1
    if 0 <= j < len(sibs):
        sibs[i], sibs[j] = sibs[j], sibs[i]
    for pos, sib in enumerate(sibs):
        sib.sort_order = pos + 1
    db.commit()
    return {"ok": True}


@router.put("/api/scenarios/{scenario_id}/fields")
def set_scenario_fields(scenario_id: int, body: LinksIn, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """Which fields this scenario (and everything below it) shows, and which of them are mandatory here."""
    s = _scenario_for(db, user, scenario_id, True)
    valid = {f.id for f in fs.form_fields(db, s.client_id, s.context, campaign_id=s.campaign_id)}
    wanted: dict[int, bool] = {}
    for link in body.links:
        fid = int(link.get("field_id", 0))
        if fid not in valid:
            raise HTTPException(422, "One of the chosen fields is not on this form")
        wanted[fid] = bool(link.get("required"))
    for row in list(db.scalars(select(ScenarioField).where(ScenarioField.scenario_id == s.id))):
        if row.field_id in wanted:
            row.required = wanted.pop(row.field_id)
        else:
            db.delete(row)
    for fid, req in wanted.items():
        db.add(ScenarioField(scenario_id=s.id, field_id=fid, required=req))
    audit(db, user, "scenario_fields_set", client_id=s.client_id, detail={"scenario": s.id})
    db.commit()
    return {"ok": True}


# ---------- try the form (setup preview; the call screen and tickets will use the same check) ----------

@router.post("/api/forms/{client_id}/{context}/check")
def check_form(client_id: int, context: str, body: CheckIn, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    _, ctx, cid = form_scope(db, user, client_id, context)
    clean, errors, names = fs.validate_answers(db, client_id, ctx, body.path, body.answers, cid)
    if errors:
        raise HTTPException(422, detail={"message": "Please fix the form", "fields": errors})
    return {"ok": True, "answers": clean, "scenario_path": names}


# ---------- value lists ----------

def _list_access(db: Session, user: User, list_id: int, write: bool) -> ValueList:
    vl = db.get(ValueList, list_id)
    if not vl:
        raise HTTPException(404, "List not found")
    perms = permissions_of(user)
    if write and "forms.manage" not in perms:
        raise HTTPException(403, "You don't have permission to change forms")
    if not write and not perms & {"forms.view", "forms.manage"}:
        # an agent reads the values of the dropdowns on the form of a client they work for
        from ..services import agent as ag
        if perms & {"calls.work", "tickets.manage"} and (vl.client_id is None or _staff(user) or ag.can_work_client(db, user, vl.client_id)):
            return vl
        raise HTTPException(403, "You don't have permission for this")
    if vl.client_id is None:
        if write and not _staff(user):
            raise HTTPException(403, "Shared lists can only be changed by DialDesk")
    elif not _staff(user) and vl.client_id != user.client_id:
        raise HTTPException(404, "List not found")
    return vl


@router.get("/api/value-lists")
def list_value_lists(client_id: Optional[int] = None, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    if not permissions_of(user) & {"forms.view", "forms.manage"}:
        raise HTTPException(403, "You don't have permission for this")
    cid = client_id if _staff(user) else user.client_id
    stmt = select(ValueList).where(ValueList.is_private.is_(False))
    stmt = stmt.where((ValueList.client_id.is_(None)) | (ValueList.client_id == cid)) if cid else stmt.where(ValueList.client_id.is_(None))
    return [fs.list_out(db, vl) for vl in db.scalars(stmt.order_by(ValueList.client_id.is_not(None), ValueList.name))]


@router.get("/api/value-lists/{list_id}")
def get_value_list(list_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    return fs.list_out(db, _list_access(db, user, list_id, False))


@router.post("/api/value-lists", status_code=201)
def create_value_list(body: ListIn, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    if "forms.manage" not in permissions_of(user):
        raise HTTPException(403, "You don't have permission to change forms")
    client_id = body.client_id if _staff(user) else user.client_id
    if client_id is None and not _staff(user):
        raise HTTPException(403, "Only DialDesk can create shared lists")
    if client_id and not db.get(Client, client_id):
        raise HTTPException(404, "Client not found")
    names = [n.strip() for n in body.level_names]
    if not all(names):
        fail({"level_names": "Give every level a name"})
    vl = ValueList(client_id=client_id, name=" ".join(body.name.split()), description=body.description, level_names=names)
    db.add(vl)
    db.flush()
    audit(db, user, "value_list_created", client_id=client_id, detail={"list": vl.id, "name": vl.name})
    db.commit()
    return fs.list_out(db, vl)


@router.put("/api/value-lists/{list_id}")
def update_value_list(list_id: int, body: ListIn, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    vl = _list_access(db, user, list_id, True)
    names = [n.strip() for n in body.level_names]
    if not all(names):
        fail({"level_names": "Give every level a name"})
    deepest = db.scalar(select(func.max(ValueItem.level)).where(ValueItem.list_id == vl.id)) or 0
    if len(names) < deepest:
        fail({"level_names": f"The list already has items on level {deepest}; remove them first"})
    vl.name, vl.description, vl.level_names, vl.is_active = " ".join(body.name.split()), body.description, names, body.is_active
    db.commit()
    return fs.list_out(db, vl)


@router.delete("/api/value-lists/{list_id}")
def delete_value_list(list_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    vl = _list_access(db, user, list_id, True)
    used = db.scalar(select(FormField.label).where(FormField.value_list_id == vl.id))
    if used:
        raise HTTPException(409, f"The field “{used}” uses this list. Change that field first.")
    db.delete(vl)
    db.commit()
    return {"ok": True}


def _item_out(db: Session, it: ValueItem) -> dict[str, Any]:
    kids = db.scalar(select(func.count()).select_from(ValueItem).where(ValueItem.parent_id == it.id))
    return {"id": it.id, "label": it.label, "level": it.level, "parent_id": it.parent_id, "is_active": it.is_active,
            "sort_order": it.sort_order, "children_count": kids}


@router.get("/api/value-lists/{list_id}/items")
def list_items(list_id: int, parent_id: Optional[int] = None, active_only: bool = False, user: User = Depends(get_current_user),
               db: Session = Depends(get_db)):
    vl = _list_access(db, user, list_id, False)
    stmt = select(ValueItem).where(ValueItem.list_id == vl.id)
    stmt = stmt.where(ValueItem.parent_id.is_(None) if parent_id is None else ValueItem.parent_id == parent_id)
    if active_only:
        stmt = stmt.where(ValueItem.is_active.is_(True))
    return [_item_out(db, it) for it in db.scalars(stmt.order_by(ValueItem.sort_order, ValueItem.label).limit(2000))]


@router.post("/api/value-lists/{list_id}/items", status_code=201)
def add_item(list_id: int, body: ItemIn, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    vl = _list_access(db, user, list_id, True)
    level, parent = 1, None
    if body.parent_id:
        parent = db.get(ValueItem, body.parent_id)
        if not parent or parent.list_id != vl.id:
            raise HTTPException(404, "Parent not found")
        level = parent.level + 1
    if level > len(vl.level_names):
        fail({"label": "This list has no deeper level. Add a level to the list first."})
    label = " ".join(body.label.split())
    dup = select(ValueItem.id).where(ValueItem.list_id == vl.id, func.lower(ValueItem.label) == label.lower(),
                                     ValueItem.parent_id.is_(None) if parent is None else ValueItem.parent_id == parent.id)
    if db.scalar(dup):
        fail({"label": "This value already exists here"})
    nxt = (db.scalar(select(func.max(ValueItem.sort_order)).where(ValueItem.list_id == vl.id)) or 0) + 1
    it = ValueItem(list_id=vl.id, parent_id=body.parent_id, level=level, label=label, sort_order=nxt, is_active=body.is_active)
    db.add(it)
    db.commit()
    return _item_out(db, it)


@router.put("/api/value-items/{item_id}")
def update_item(item_id: int, body: ItemIn, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    it = db.get(ValueItem, item_id)
    if not it:
        raise HTTPException(404, "Value not found")
    _list_access(db, user, it.list_id, True)
    label = " ".join(body.label.split())
    dup = select(ValueItem.id).where(ValueItem.list_id == it.list_id, func.lower(ValueItem.label) == label.lower(),
                                     ValueItem.id != it.id,
                                     ValueItem.parent_id.is_(None) if it.parent_id is None else ValueItem.parent_id == it.parent_id)
    if db.scalar(dup):
        fail({"label": "This value already exists here"})
    it.label, it.is_active = label, body.is_active
    db.commit()
    return _item_out(db, it)


@router.delete("/api/value-items/{item_id}")
def delete_item(item_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    it = db.get(ValueItem, item_id)
    if not it:
        raise HTTPException(404, "Value not found")
    _list_access(db, user, it.list_id, True)
    db.delete(it)  # values below it go too
    db.commit()
    return {"ok": True}


@router.post("/api/value-lists/{list_id}/import")
async def import_list(list_id: int, file: UploadFile = File(...), header: bool = Form(True), user: User = Depends(get_current_user),
                      db: Session = Depends(get_db)):
    """CSV or Excel with one column per level, for example Country, State, City."""
    vl = _list_access(db, user, list_id, True)
    name = (file.filename or "").lower()
    content = await file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(422, "The file must be smaller than 5 MB")
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = [["" if c is None else str(c) for c in r] for r in wb.active.iter_rows(values_only=True)]
    elif name.endswith((".csv", ".txt")):
        rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig", errors="replace"))))
    else:
        raise HTTPException(422, "Upload a .csv or .xlsx file")
    rows = [r for r in rows if any(c.strip() for c in r)]
    if header and rows:
        rows = rows[1:]
    if len(rows) > 50_000:
        raise HTTPException(422, "Too many rows; split the file into parts of 50,000")
    added, skipped = fs.import_rows(db, vl, rows)
    audit(db, user, "value_list_imported", client_id=vl.client_id, detail={"list": vl.id, "added": added})
    db.commit()
    return {"added": added, "skipped": skipped, "rows": len(rows)}


# ---------- copy a form ----------

@router.post("/api/forms/{client_id}/{context}/copy-from")
def copy_from(client_id: int, context: str, body: CopyIn, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """Copies fields, scenarios and lists into this form.
    Campaign form: from another campaign (a client user may copy between their own company's campaigns).
    In-call form: from another client (DialDesk staff only, as it reads another company's setup)."""
    target, ctx, t_cid = form_scope(db, user, client_id, context, write=True)
    source = db.get(Client, body.source_client_id)
    if not source:
        raise HTTPException(422, "Choose a client to copy from")
    if ctx == "out_call":
        src_campaign = db.get(Campaign, body.source_campaign_id) if body.source_campaign_id else None
        if not src_campaign or src_campaign.client_id != source.id:
            raise HTTPException(422, "Choose the campaign to copy from")
        if src_campaign.id == t_cid:
            raise HTTPException(422, "Choose a different campaign to copy from")
        if not _staff(user) and (source.id != user.client_id or target.id != user.client_id):
            raise HTTPException(403, "You can only copy between your own company's campaigns")
        s_cid = src_campaign.id
    else:
        if not _staff(user):
            raise HTTPException(403, "Only DialDesk staff can copy a form from another client")
        if source.id == target.id:
            raise HTTPException(422, "Choose a different client to copy from")
        s_cid = None

    have = db.scalar(select(func.count()).select_from(FormField).where(
        FormField.client_id == target.id, FormField.context == ctx, fs.in_scope(FormField, t_cid))) \
        or db.scalar(select(func.count()).select_from(Scenario).where(
            Scenario.client_id == target.id, Scenario.context == ctx, fs.in_scope(Scenario, t_cid)))
    if have and not body.replace:
        raise HTTPException(409, "This form already has content. Confirm to replace it.")
    if have:
        for f in fs.form_fields(db, target.id, ctx, campaign_id=t_cid):
            if f.value_list_id:
                vl = db.get(ValueList, f.value_list_id)
                if vl and vl.is_private:
                    db.delete(vl)
        db.query(FormField).filter(FormField.client_id == target.id, FormField.context == ctx,
                                   fs.in_scope(FormField, t_cid)).delete(synchronize_session=False)
        db.query(Scenario).filter(Scenario.client_id == target.id, Scenario.context == ctx,
                                  fs.in_scope(Scenario, t_cid)).delete(synchronize_session=False)
        db.flush()

    list_map: dict[int, int] = {}

    def copy_list(src: ValueList) -> int:
        if src.id in list_map:
            return list_map[src.id]
        # shared lists, and a client's own lists when copying inside the same client, are simply used again;
        # a field's private values are always copied so each form keeps its own
        if src.client_id is None or (src.client_id == target.id and not src.is_private):
            list_map[src.id] = src.id
            return src.id
        new = ValueList(client_id=target.id, name=src.name, description=src.description, level_names=list(src.level_names),
                        is_private=src.is_private, is_active=src.is_active)
        db.add(new)
        db.flush()
        item_map: dict[int, int] = {}
        for it in db.scalars(select(ValueItem).where(ValueItem.list_id == src.id).order_by(ValueItem.level, ValueItem.id)):
            c = ValueItem(list_id=new.id, parent_id=item_map.get(it.parent_id) if it.parent_id else None, level=it.level,
                          label=it.label, sort_order=it.sort_order, is_active=it.is_active)
            db.add(c)
            db.flush()
            item_map[it.id] = c.id
        list_map[src.id] = new.id
        return new.id

    field_map: dict[int, int] = {}
    src_fields = fs.form_fields(db, source.id, ctx, campaign_id=s_cid)
    for f in src_fields:
        lid = copy_list(db.get(ValueList, f.value_list_id)) if f.value_list_id else None
        n = FormField(client_id=target.id, context=ctx, campaign_id=t_cid, key=f.key, label=f.label,
                      field_type=f.field_type, validation=f.validation, rules=dict(f.rules or {}), required=f.required,
                      help_text=f.help_text, sort_order=f.sort_order, is_active=f.is_active, show_always=f.show_always,
                      value_list_id=lid, list_level=f.list_level)
        db.add(n)
        db.flush()
        field_map[f.id] = n.id
    for f in src_fields:  # parents are linked once every field exists
        if f.parent_field_id:
            db.get(FormField, field_map[f.id]).parent_field_id = field_map.get(f.parent_field_id)

    scen_map: dict[int, int] = {}
    src_scen = list(db.scalars(select(Scenario).where(Scenario.client_id == source.id, Scenario.context == ctx,
                                                      fs.in_scope(Scenario, s_cid)).order_by(Scenario.level, Scenario.sort_order)))
    for sc in src_scen:
        n = Scenario(client_id=target.id, context=ctx, campaign_id=t_cid, parent_id=scen_map.get(sc.parent_id) if sc.parent_id else None,
                     level=sc.level, name=sc.name, sort_order=sc.sort_order, outcome=sc.outcome, is_active=sc.is_active)
        db.add(n)
        db.flush()
        scen_map[sc.id] = n.id
        for link in db.scalars(select(ScenarioField).where(ScenarioField.scenario_id == sc.id)):
            if link.field_id in field_map:
                db.add(ScenarioField(scenario_id=n.id, field_id=field_map[link.field_id], required=link.required))
    audit(db, user, "form_copied", client_id=target.id, detail={"from_client": source.id, "from_campaign": s_cid, "context": ctx})
    db.commit()
    return {"fields": len(field_map), "scenarios": len(scen_map)}
