"""Campaign data: reading uploads, the calling queue, and keeping a lead up to date after each call."""
import csv
import io
import re
from datetime import datetime, timedelta, timezone
from typing import Any, Optional

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from ..models import Customer, Interaction, Lead, User
from .phone import normalize_phone, valid_phone

MAX_ROWS = 50000
MAX_BYTES = 10 * 1024 * 1024
MAX_ATTEMPTS = 3            # a lead that never connects is retired after this many tries
LOCK_MINUTES = 120          # a lead taken by an agent stays theirs this long
RETRY_AFTER = timedelta(hours=2)
CALLBACK_AFTER = timedelta(hours=4)
SKIP_RETRY = timedelta(minutes=15)

FIELDS = ("name", "mobile", "email", "city", "ext_ref")
GUESS = {
    "name": ("name", "customer name", "full name", "customer", "contact name", "client name", "first name", "firstname"),
    "name2": ("last name", "lastname", "surname"),  # optional: added after the name (a dialer export has first and last name apart)
    "mobile": ("mobile", "mobile number", "mobile no", "phone", "phone number", "contact number", "contact no", "number", "msisdn"),
    "email": ("email", "email id", "e mail", "mail"),
    "city": ("city", "town", "location"),
    "ext_ref": ("ref", "reference", "ref no", "policy no", "policy number", "customer id", "account no", "id", "order no"),
}


def _norm(h: str) -> str:
    return re.sub(r"[^a-z0-9 ]", " ", str(h or "").lower()).strip()


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_file(filename: str, raw: bytes) -> tuple[list[str], list[list[str]]]:
    """Header row and data rows of a CSV or Excel file, as text."""
    if len(raw) > MAX_BYTES:
        raise HTTPException(422, f"The file is larger than {MAX_BYTES // (1024 * 1024)} MB. Split it into smaller files.")
    name = (filename or "").lower()
    rows: list[list[str]] = []
    try:
        if name.endswith((".xlsx", ".xlsm")):
            ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
            for r in ws.iter_rows(values_only=True):
                rows.append([_cell(v) for v in r])
                if len(rows) > MAX_ROWS + 1:
                    break
        elif name.endswith((".csv", ".txt")):
            text = raw.decode("utf-8-sig", errors="replace")
            for r in csv.reader(io.StringIO(text)):
                rows.append([c.strip() for c in r])
                if len(rows) > MAX_ROWS + 1:
                    break
        else:
            raise HTTPException(422, "Upload a .csv or .xlsx file.")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, "This file could not be read. Save it as .csv or .xlsx and try again.")
    rows = [r for r in rows if any(c for c in r)]
    if len(rows) < 2:
        raise HTTPException(422, "The file needs a header row and at least one data row.")
    if len(rows) - 1 > MAX_ROWS:
        raise HTTPException(422, f"The file has more than {MAX_ROWS} rows. Split it into smaller files.")
    header = [h or f"Column {i + 1}" for i, h in enumerate(rows[0])]
    width = len(header)
    return header, [(r + [""] * width)[:width] for r in rows[1:]]


def guess_mapping(header: list[str]) -> dict[str, Optional[str]]:
    mapping: dict[str, Optional[str]] = {}
    used: set[str] = set()
    for field, names in GUESS.items():
        found = next((h for h in header if h not in used and _norm(h) in names), None)
        mapping[field] = found
        if found:
            used.add(found)
    return mapping


def build_leads(header: list[str], rows: list[list[str]], mapping: dict[str, Any], existing: set[str]):
    """Turns rows into lead dicts. Returns (good, duplicates, errors); errors are (row number, reason)."""
    idx = {h: i for i, h in enumerate(header)}
    for f in ("name", "mobile"):
        if mapping.get(f) not in idx:
            raise HTTPException(422, detail={"message": "Choose the columns for name and mobile number.", "fields": {f: "Required"}})
    pick = {f: idx[mapping[f]] for f in FIELDS + ("name2",) if mapping.get(f) in idx}
    mapped = {mapping.get(f) for f in FIELDS + ("name2",)}
    extra_cols = [c for c in (mapping.get("extra") or []) if c in idx and c not in mapped]
    good: list[dict[str, Any]] = []
    errors: list[tuple[int, str]] = []
    dups = 0
    seen = set(existing)
    for n, r in enumerate(rows, start=2):
        last = r[pick["name2"]].strip() if "name2" in pick else ""
        name = " ".join(f"{r[pick['name']]} {'' if last in ('.', '-', 'N/A', 'n/a') else last}".split())
        mobile = normalize_phone(r[pick["mobile"]])
        if len(name) < 2:
            errors.append((n, "Name is missing"))
            continue
        if not valid_phone(mobile):
            errors.append((n, f"Mobile number {r[pick['mobile']][:20]!r} is not valid"))
            continue
        if mobile in seen:
            dups += 1
            continue
        seen.add(mobile)
        email = r[pick["email"]].strip().lower() if "email" in pick else ""
        if email and ("@" not in email or "." not in email.split("@")[-1]):
            email = ""
        good.append({"name": name[:150], "mobile": mobile, "email": email[:255] or None,
                     "city": (r[pick["city"]].strip()[:100] or None) if "city" in pick else None,
                     "ext_ref": (r[pick["ext_ref"]].strip()[:100] or None) if "ext_ref" in pick else None,
                     "extra": {c: r[idx[c]][:200] for c in extra_cols if r[idx[c]]}})
    return good, dups, errors


# ---------- the queue ----------

def _now() -> datetime:
    return datetime.now(timezone.utc)


def release_stale(db: Session, campaign_id: int) -> None:
    for l in db.scalars(select(Lead).where(Lead.campaign_id == campaign_id, Lead.locked_by_id.is_not(None), Lead.locked_until < _now())):
        l.locked_by_id = l.locked_until = None


def available(user_id: int, campaign_id: int, batch_id: Optional[int] = None):
    """Leads this agent may take now: open, not locked by someone else, due, and theirs or in the shared queue.
    batch_id narrows it to one allocation (an upload); 0 means the data that came in through the client's API."""
    now = _now()
    batch = [] if batch_id is None else [Lead.batch_id.is_(None) if batch_id == 0 else Lead.batch_id == batch_id]
    return select(Lead).where(*batch,
        Lead.campaign_id == campaign_id, Lead.status.in_(("new", "callback")),
        or_(Lead.next_call_at.is_(None), Lead.next_call_at <= now),
        or_(Lead.locked_by_id.is_(None), Lead.locked_by_id == user_id, Lead.locked_until < now),
        or_(Lead.assigned_to_id.is_(None), Lead.assigned_to_id == user_id))


def pick_next(db: Session, user_id: int, campaign_id: int, batch_id: Optional[int] = None) -> Optional[Lead]:
    """Due call-backs first, then the agent's own data, then the shared queue (oldest first)."""
    rows = list(db.scalars(available(user_id, campaign_id, batch_id).order_by(Lead.id).limit(300).with_for_update(skip_locked=True)))
    if not rows:
        return None
    floor = datetime.min.replace(tzinfo=timezone.utc)
    rows.sort(key=lambda l: (l.status != "callback", l.assigned_to_id != user_id, l.next_call_at or floor, l.id))
    return rows[0]


def _customer_for(db: Session, user: User, lead: Lead) -> Customer:
    cust = db.scalar(select(Customer).where(Customer.client_id == lead.client_id, Customer.mobile == lead.mobile))
    if not cust:
        cust = Customer(client_id=lead.client_id, name=lead.name, mobile=lead.mobile, email=lead.email, city=lead.city,
                        ext_ref=lead.ext_ref, created_by_id=user.id)
        db.add(cust)
        db.flush()
    return cust


def start_call(db: Session, user: User, lead: Lead) -> Interaction:
    """Locks the lead to the agent and opens a manual call for it."""
    now = _now()
    lead.locked_by_id, lead.locked_until = user.id, now + timedelta(minutes=LOCK_MINUTES)
    lead.attempts += 1
    lead.last_called_at = now
    cust = _customer_for(db, user, lead)
    it = Interaction(client_id=lead.client_id, campaign_id=lead.campaign_id, customer_id=cust.id, agent_id=user.id, mode="manual",
                     phone=lead.mobile)
    db.add(it)
    db.flush()
    lead.interaction_id = it.id
    return it


def attach_dialed(db: Session, user: User, it: Interaction) -> None:
    """A predictive call came in for a number that is in the campaign data: connect the two."""
    if not it.campaign_id:
        return
    lead = db.scalar(select(Lead).where(Lead.campaign_id == it.campaign_id, Lead.mobile == it.phone))
    if not lead:
        return
    now = _now()
    lead.interaction_id, lead.last_called_at = it.id, now
    lead.attempts += 1
    lead.locked_by_id, lead.locked_until = user.id, now + timedelta(minutes=LOCK_MINUTES)
    if not it.customer_id:
        it.customer_id = _customer_for(db, user, lead).id


def lead_of(db: Session, it: Interaction) -> Optional[Lead]:
    if not it.campaign_id:
        return None
    return db.scalar(select(Lead).where(Lead.campaign_id == it.campaign_id, Lead.interaction_id == it.id))


def after_tag(db: Session, it: Interaction, callback_at: Optional[datetime]) -> None:
    lead = lead_of(db, it)
    if not lead:
        return
    now = _now()
    lead.locked_by_id = lead.locked_until = None
    lead.last_outcome = it.outcome
    if it.outcome == "dnc":
        lead.status, lead.next_call_at = "dnc", None
    elif it.outcome == "callback":
        lead.status = "callback"
        lead.next_call_at = callback_at if callback_at and callback_at > now else now + CALLBACK_AFTER
        lead.assigned_to_id = it.agent_id  # the agent who spoke to the customer follows up
    elif it.outcome == "not_connected":
        if lead.attempts >= MAX_ATTEMPTS:
            lead.status, lead.next_call_at = "done", None
        else:
            lead.status, lead.next_call_at = "new", now + RETRY_AFTER
    else:  # connected, converted, closed, or no outcome set on the scenario
        lead.status, lead.next_call_at = "done", None


def after_skip(db: Session, it: Interaction) -> None:
    lead = lead_of(db, it)
    if lead:
        lead.locked_by_id = lead.locked_until = None
        lead.next_call_at = _now() + SKIP_RETRY


def lead_out(l: Lead, agents: Optional[dict[int, str]] = None) -> dict[str, Any]:
    return {"id": l.id, "name": l.name, "mobile": l.mobile, "email": l.email, "city": l.city, "ext_ref": l.ext_ref,
            "extra": l.extra or {}, "status": l.status, "attempts": l.attempts, "assigned_to_id": l.assigned_to_id,
            "assigned_to": (agents or {}).get(l.assigned_to_id) if l.assigned_to_id else None,
            "next_call_at": l.next_call_at, "last_outcome": l.last_outcome, "last_called_at": l.last_called_at,
            "batch_id": l.batch_id}


# ---------- records posted by a client (webhook) ----------

ALIASES = {
    "name": ("name", "customer_name", "full_name", "customername", "fullname", "contact_name"),
    "mobile": ("mobile", "mobile_number", "mobileno", "phone", "phone_number", "phonenumber", "number", "msisdn", "contact_number"),
    "email": ("email", "email_id", "emailid", "mail"),
    "city": ("city", "town"),
    "ext_ref": ("ref", "reference", "ext_ref", "policy_no", "policy_number", "customer_id", "account_no", "order_no", "id"),
}


def build_from_records(records: list[Any], existing: set[str]):
    """Like build_leads, for JSON records. Returns (good, results) where results has one entry per record."""
    good: list[dict[str, Any]] = []
    results: list[dict[str, Any]] = []
    seen = set(existing)
    for i, rec in enumerate(records):
        if not isinstance(rec, dict):
            results.append({"index": i, "status": "invalid", "reason": "Each record must be an object"})
            continue
        low = {str(k).strip().lower().replace(" ", "_"): v for k, v in rec.items()}
        pick: dict[str, Any] = {}
        used: set[str] = set()
        for field, names in ALIASES.items():
            for n in names:
                if n in low and low[n] not in (None, ""):
                    pick[field] = str(low[n]).strip()
                    used.add(n)
                    break
        name = " ".join(pick.get("name", "").split())
        mobile = normalize_phone(pick.get("mobile", ""))
        if len(name) < 2:
            results.append({"index": i, "status": "invalid", "reason": "name is required"})
            continue
        if not valid_phone(mobile):
            results.append({"index": i, "status": "invalid", "reason": "mobile is required and must be a valid number"})
            continue
        if mobile in seen:
            results.append({"index": i, "status": "duplicate", "mobile": mobile})
            continue
        seen.add(mobile)
        email = pick.get("email", "").lower()
        if email and ("@" not in email or "." not in email.split("@")[-1]):
            email = ""
        extra = {str(k)[:60]: str(v)[:200] for k, v in low.items() if k not in used and v not in (None, "") and not isinstance(v, (dict, list))}
        good.append({"name": name[:150], "mobile": mobile, "email": email[:255] or None, "city": pick.get("city", "")[:100] or None,
                     "ext_ref": pick.get("ext_ref", "")[:100] or None, "extra": dict(list(extra.items())[:20])})
        results.append({"index": i, "status": "accepted", "mobile": mobile})
    return good, results
